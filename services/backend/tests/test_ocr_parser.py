"""Tests for Task B15: Local OCR and DOCX rendering fallbacks."""
import io
import pytest
from PIL import Image, ImageDraw
import docx
from pypdf import PdfWriter

from app.services.parser import (
    DocumentParsingError,
    parse_docx_bytes,
    parse_document_content,
    parse_pdf_bytes,
    run_ocr_on_image_bytes,
)


def create_image_pdf(text: str) -> bytes:
    """Create a 1-page PDF containing a rendered image with the given text."""
    img = Image.new("RGB", (400, 100), color=(255, 255, 255))
    d = ImageDraw.Draw(img)
    d.text((10, 30), text, fill=(0, 0, 0))

    pdf_buf = io.BytesIO()
    img.save(pdf_buf, format="PDF")
    return pdf_buf.getvalue()


def create_blank_image_pdf() -> bytes:
    """Create a 1-page PDF containing a blank white image (no text)."""
    img = Image.new("RGB", (200, 100), color=(255, 255, 255))
    pdf_buf = io.BytesIO()
    img.save(pdf_buf, format="PDF")
    return pdf_buf.getvalue()


def create_test_docx(paragraphs: list[str], table_data: list[list[str]] = None) -> bytes:
    """Create a DOCX file with paragraphs and optional table."""
    doc = docx.Document()
    for p in paragraphs:
        doc.add_paragraph(p)

    if table_data:
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        for r_idx, row in enumerate(table_data):
            for c_idx, cell_value in enumerate(row):
                table.cell(r_idx, c_idx).text = cell_value

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_ocr_recovers_scanned_pdf():
    """Scanned PDF without text layer triggers local OCR and preserves exact span provenance."""
    pdf_bytes = create_image_pdf("Senior Backend Engineer Go")
    res = parse_pdf_bytes(pdf_bytes)

    assert res.quality_report["ocr_applied"] is True
    assert 1 in res.quality_report["ocr_pages"]
    assert "Senior" in res.canonical_text or "Backend" in res.canonical_text or "Go" in res.canonical_text

    # Verify exact codepoint alignment invariant on OCR text
    assert len(res.spans) > 0
    for s in res.spans:
        assert res.canonical_text[s.start_cp : s.end_cp] == s.text
        assert s.span_id.startswith("spn_")


@pytest.mark.asyncio
async def test_empty_scanned_pdf_flags_warning():
    """Purely blank scanned PDF produces quality report with suspected_scanned warning."""
    blank_pdf = create_blank_image_pdf()
    res = parse_pdf_bytes(blank_pdf)
    assert res.canonical_text == ""
    assert res.quality_report["suspected_scanned"] is True
    assert res.quality_report["empty_pages"] == 1


@pytest.mark.asyncio
async def test_docx_table_and_columns_parsing():
    """DOCX parsing extracts paragraphs and table cell contents separated by pipe delimiter."""
    docx_bytes = create_test_docx(
        paragraphs=["Kinh nghiệm làm việc", "Kỹ sư phần mềm cấp cao tại Công ty ABC"],
        table_data=[
            ["Kỹ năng", "Mức độ"],
            ["Golang", "Chuyên sâu"],
            ["PostgreSQL", "Nâng cao"],
        ],
    )

    res = parse_docx_bytes(docx_bytes)
    assert "Kinh nghiệm làm việc" in res.canonical_text
    assert "Golang | Chuyên sâu" in res.canonical_text
    assert "PostgreSQL | Nâng cao" in res.canonical_text
    assert res.quality_report["rendered_page_count"] >= 1
    assert res.quality_report["docx_pages_verified"] is True
    assert res.quality_report["excessive_page_count"] is False

    # Codepoint invariant check
    for s in res.spans:
        assert res.canonical_text[s.start_cp : s.end_cp] == s.text


def test_docx_uses_rendered_pages_for_multpage_provenance():
    doc = docx.Document()
    doc.add_paragraph("First page backend Python experience.")
    doc.add_page_break()
    doc.add_paragraph("Second page database PostgreSQL project.")
    buf = io.BytesIO()
    doc.save(buf)
    result = parse_docx_bytes(buf.getvalue())
    assert result.page_count == 2
    assert result.quality_report["docx_pages_verified"] is True
    assert {span.page_number for span in result.spans} == {1, 2}
    assert all(result.canonical_text[span.start_cp:span.end_cp] == span.text for span in result.spans)


def test_docx_missing_renderer_is_technical_error(monkeypatch):
    from app.services import parser
    monkeypatch.setattr(parser.shutil, "which", lambda _: None)
    with pytest.raises(DocumentParsingError) as exc:
        parse_docx_bytes(create_test_docx(["Backend experience"]))
    assert exc.value.error_code == "DOCX_RENDERER_UNAVAILABLE"


@pytest.mark.asyncio
async def test_empty_docx_returns_empty_result():
    """Empty DOCX with no paragraphs or tables returns empty canonical text and 1 empty page."""
    empty_docx = create_test_docx([])
    res = parse_docx_bytes(empty_docx)
    assert res.canonical_text == ""
    assert res.quality_report["empty_pages"] == 1


@pytest.mark.asyncio
async def test_unsupported_mime_raises_error():
    """Unsupported MIME type raises UNSUPPORTED_MIME."""
    with pytest.raises(DocumentParsingError) as exc_info:
        parse_document_content(b"test", "image/png")

    assert exc_info.value.error_code == "UNSUPPORTED_MIME"


@pytest.mark.asyncio
async def test_run_ocr_graceful_on_invalid_image():
    """run_ocr_on_image_bytes returns empty string on corrupt image bytes without crashing."""
    corrupt_bytes = b"not_an_image_binary_data"
    result = run_ocr_on_image_bytes(corrupt_bytes)
    assert result == ""
