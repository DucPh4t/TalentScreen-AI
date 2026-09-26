"""Comprehensive tests for Task B13: HR Revisions, Attestation, and Attested Final Decision Flow."""
from __future__ import annotations

from datetime import datetime, timedelta, timezone
from decimal import Decimal
import io
import json
import uuid
import docx
import pytest
from httpx import ASGITransport, AsyncClient
from sqlalchemy import delete, select

from app.db.models import (
    Application,
    AssessmentRun,
    BudgetReservation,
    Candidate,
    CriterionAssessment,
    CriterionEvidence,
    Decision,
    Document,
    HRRevision,
    JDVersion,
    Job,
    RawAccessGrant,
    Requisition,
    RequisitionMembership,
    ReviewAttestation,
    RubricCriterion,
    RubricVersion,
    SanitizedVersion,
    SourceSpan,
    User,
    UserAccountRole,
)
from app.domain.authorization import SESSION_COOKIE_NAME
from app.domain.enums import (
    AccountRole,
    CriterionId,
    CriterionOutcome,
    DecisionBasis,
    DecisionOutcome,
    DocumentSafetyStatus,
    MembershipRole,
    Recommendation,
    RequisitionStatus,
    RubricStatus,
    SanitizedVersionStatus,
    UserStatus,
)
from app.domain.security import hash_password
from app.main import app
from app.services.auth import create_session


@pytest.fixture(autouse=True)
async def cleanup_decision_tables(test_session_factory):
    """Ensure clean slate across tests."""
    async with test_session_factory() as session:
        await session.execute(delete(Decision))
        await session.execute(delete(ReviewAttestation))
        await session.execute(delete(HRRevision))
        await session.execute(delete(CriterionEvidence))
        await session.execute(delete(CriterionAssessment))
        await session.execute(delete(AssessmentRun))
        await session.execute(delete(BudgetReservation))
        await session.execute(delete(Job))
        await session.execute(delete(SourceSpan))
        await session.commit()
    yield
    async with test_session_factory() as session:
        await session.execute(delete(Decision))
        await session.execute(delete(ReviewAttestation))
        await session.execute(delete(HRRevision))
        await session.execute(delete(CriterionEvidence))
        await session.execute(delete(CriterionAssessment))
        await session.execute(delete(AssessmentRun))
        await session.execute(delete(BudgetReservation))
        await session.execute(delete(Job))
        await session.execute(delete(SourceSpan))
        await session.commit()


@pytest.fixture
def sample_docx_cv() -> bytes:
    doc = docx.Document()
    doc.add_paragraph("Kinh nghiệm làm việc chuyên sâu Backend Python và FastAPI.")
    doc.add_paragraph("Thiết kế kiến trúc hệ thống microservices chịu tải cao và tối ưu cơ sở dữ liệu PostgreSQL.")
    doc.add_paragraph("Triển khai hạ tầng container với Docker và Kubernetes.")
    doc.add_paragraph("Kỹ năng giao tiếp và làm việc nhóm hiệu quả.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def setup_test_context(session, sample_docx_cv):
    """Setup org, owner, reviewer, requisition, approved rubric, candidate, application, and approved sanitized version."""
    from app.services.requisition import get_or_create_default_org
    org = await get_or_create_default_org(session)

    # 1. Users
    owner = User(
        id=uuid.uuid4(),
        login_name=f"owner_{uuid.uuid4().hex[:8]}",
        display_name="Owner HR",
        password_hash=hash_password("Password123!"),
        status=UserStatus.ACTIVE,
    )
    reviewer = User(
        id=uuid.uuid4(),
        login_name=f"rev_{uuid.uuid4().hex[:8]}",
        display_name="Reviewer HR",
        password_hash=hash_password("Password123!"),
        status=UserStatus.ACTIVE,
    )
    session.add_all([owner, reviewer])
    await session.flush()
    session.add(UserAccountRole(user_id=owner.id, role=AccountRole.RECRUITER))
    session.add(UserAccountRole(user_id=reviewer.id, role=AccountRole.RECRUITER))

    o_token, o_csrf, _ = await create_session(session, owner)
    r_token, r_csrf, _ = await create_session(session, reviewer)

    # 2. Requisition & Memberships
    req = Requisition(
        id=uuid.uuid4(),
        organization_id=org.id,
        title="Senior Python Engineer",
        status=RequisitionStatus.OPEN,
        row_version=1,
    )
    session.add(req)
    await session.flush()
    session.add(RequisitionMembership(requisition_id=req.id, user_id=owner.id, membership_role=MembershipRole.OWNER))
    session.add(RequisitionMembership(requisition_id=req.id, user_id=reviewer.id, membership_role=MembershipRole.REVIEWER))
    await session.flush()

    # 3. JD & Rubric
    jd = JDVersion(
        id=uuid.uuid4(),
        requisition_id=req.id,
        version_no=1,
        source_text="Senior Python Backend Engineer JD",
        text_hash="a" * 64,
        created_by=owner.id,
    )
    session.add(jd)
    await session.flush()
    req.current_jd_version_id = jd.id

    rubric = RubricVersion(
        id=uuid.uuid4(),
        requisition_id=req.id,
        jd_version_id=jd.id,
        version_no=1,
        status=RubricStatus.APPROVED,
        content_hash="b" * 64,
        approved_by=owner.id,
        approved_at=datetime.now(timezone.utc),
    )
    session.add(rubric)
    await session.flush()

    for cid in CriterionId:
        crit = RubricCriterion(
            rubric_version_id=rubric.id,
            criterion_id=cid.value,
            label_vi=cid.name,
            weight=20 if cid == CriterionId.PYTHON_BACKEND else 16,
            description_vi=f"Description for {cid.name}",
            anchors={"0": "None", "1": "Basic", "2": "Intermediate", "3": "Advanced", "4": "Expert"},
        )
        session.add(crit)

    # 4. Candidate & Application
    cand = Candidate(id=uuid.uuid4(), organization_id=org.id, public_label=f"CAND-{uuid.uuid4().hex[:6].upper()}", status="active")
    session.add(cand)
    await session.flush()

    app_obj = Application(
        id=uuid.uuid4(),
        requisition_id=req.id,
        candidate_id=cand.id,
        status="active",
        generation=1,
        row_version=1,
    )
    session.add(app_obj)
    await session.flush()

    # 5. Document & Sanitized Version & Spans
    doc = Document(
        id=uuid.uuid4(),
        application_id=app_obj.id,
        version_no=1,
        kind="cv",
        original_name_private="cv.docx",
        mime_verified="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        byte_size=len(sample_docx_cv),
        sha256="c" * 64,
        blob_key="docs/test.docx",
        ingestion_status="parsed",
        safety_status=DocumentSafetyStatus.PASSED,
        created_by=owner.id,
    )
    session.add(doc)
    await session.flush()

    app_obj.current_document_id = doc.id

    sanitized = SanitizedVersion(
        id=uuid.uuid4(),
        application_id=app_obj.id,
        document_id=doc.id,
        version_no=1,
        status=SanitizedVersionStatus.APPROVED,
        canonical_text="Kinh nghiệm làm việc chuyên sâu Backend Python và FastAPI.\nThiết kế microservices PostgreSQL.",
        sha256="d" * 64,
        approved_by=owner.id,
        approved_at=datetime.now(timezone.utc),
    )
    session.add(sanitized)
    await session.flush()

    app_obj.current_sanitized_version_id = sanitized.id

    span_id = f"spn_{uuid.uuid4().hex[:24]}"
    span1 = SourceSpan(
        span_id=span_id,
        full_hash="e" * 64,
        sanitized_version_id=sanitized.id,
        start_cp=0,
        end_cp=58,
        page_number=1,
        language="vi",
        text="Kinh nghiệm làm việc chuyên sâu Backend Python và FastAPI.",
    )
    session.add(span1)
    await session.commit()

    return {
        "owner": owner,
        "o_token": o_token,
        "o_csrf": o_csrf,
        "reviewer": reviewer,
        "r_token": r_token,
        "r_csrf": r_csrf,
        "req_id": req.id,
        "rubric_id": rubric.id,
        "cand_id": cand.id,
        "app_id": app_obj.id,
        "doc_id": doc.id,
        "sanitized_id": sanitized.id,
        "span_id": span1.span_id,
        "span_text": span1.text,
    }


# ---------------- Test Cases ---------------- #


@pytest.mark.asyncio
async def test_closed_requisition_keeps_review_history_readable(test_session_factory, sample_docx_cv):
    async with test_session_factory() as session:
        ctx = await setup_test_context(session, sample_docx_cv)
        req = (await session.execute(select(Requisition).where(Requisition.id == ctx["req_id"]))).scalar_one()
        req.status = RequisitionStatus.CLOSED
        await session.commit()

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test", cookies={SESSION_COOKIE_NAME: ctx["o_token"]}, headers={"X-CSRF-Token": ctx["o_csrf"]}) as client:
        decisions = await client.get(f"/api/v1/applications/{ctx['app_id']}/decisions")
        revisions = await client.get(f"/api/v1/applications/{ctx['app_id']}/hr-revisions")
        assert decisions.status_code == 200
        assert decisions.json() == []
        assert revisions.status_code == 200


@pytest.mark.asyncio
async def test_create_and_update_hr_revision_with_recalculated_scores(test_session_factory, sample_docx_cv):
    """Reviewer creates draft HR revision, system calculates scores server-side, reviewer updates draft."""
    async with test_session_factory() as session:
        ctx = await setup_test_context(session, sample_docx_cv)

    transport = ASGITransport(app=app)
    cookies = {SESSION_COOKIE_NAME: ctx["r_token"]}
    headers = {"X-CSRF-Token": ctx["r_csrf"]}

    async with AsyncClient(transport=transport, base_url="http://test", cookies=cookies, headers=headers) as client:
        # 1. Create draft HR revision
        criteria = [
            {
                "criterion_id": cid.value,
                "status": "assessed",
                "score": 3,
                "evidence": [{"span_id": ctx["span_id"], "quote": ctx["span_text"]}],
                "rationale": f"Reviewer assessed competency in {cid.value} as level 3.",
                "missing_information": [],
            }
            for cid in CriterionId
        ]
        change_reasons = {
            "python_backend": {
                "reason_code": "missed_evidence",
                "note": "Candidate mentioned FastAPI microservices in project section which AI missed.",
            }
        }

        res = await client.post(
            f"/api/v1/applications/{ctx['app_id']}/hr-revisions",
            json={
                "expected_application_version": 1,
                "source_snapshot_ref": {
                    "document_id": str(ctx["doc_id"]),
                    "sanitized_version_id": str(ctx["sanitized_id"]),
                    "rubric_version_id": str(ctx["rubric_id"]),
                    "application_generation": 1,
                },
                "criteria": criteria,
                "change_reasons": change_reasons,
                "summary_reason": "Reviewer manual revision after deep CV check.",
                "proposed_decision": "advance",
            },
        )
        assert res.status_code == 201
        data = res.json()
        assert data["status"] == "draft"
        assert Decimal(str(data["comparable_score"])) == Decimal("75.00")
        assert Decimal(str(data["coverage"])) == Decimal("1.00")
        assert data["recommendation"] == "consider_next_round"
        rev_id = data["id"]

        # 2. Get revision detail
        get_res = await client.get(f"/api/v1/hr-revisions/{rev_id}")
        assert get_res.status_code == 200
        assert get_res.json()["is_stale"] is False

        # 3. Update draft revision
        # Change python_backend score to 4
        criteria[0]["score"] = 4
        up_res = await client.put(
            f"/api/v1/hr-revisions/{rev_id}",
            json={
                "expected_revision_version": 1,
                "criteria": criteria,
                "change_reasons": change_reasons,
                "summary_reason": "Upgraded python_backend to level 4 due to extensive experience.",
                "proposed_decision": "advance",
            },
        )
        assert up_res.status_code == 200
        up_data = up_res.json()
        # New comparable score: 20*(4/4) + 16*5*(3/4) = 20 + 60 = 80.00
        assert Decimal(str(up_data["comparable_score"])) == Decimal("80.00")


@pytest.mark.asyncio
async def test_anti_discrimination_scan_on_hr_revision(test_session_factory, sample_docx_cv):
    """Attempting to include age, gender, or marital status in HR notes triggers 422."""
    async with test_session_factory() as session:
        ctx = await setup_test_context(session, sample_docx_cv)

    transport = ASGITransport(app=app)
    cookies = {SESSION_COOKIE_NAME: ctx["r_token"]}
    headers = {"X-CSRF-Token": ctx["r_csrf"]}

    async with AsyncClient(transport=transport, base_url="http://test", cookies=cookies, headers=headers) as client:
        criteria = [
            {
                "criterion_id": cid.value,
                "status": "assessed",
                "score": 3,
                "evidence": [{"span_id": ctx["span_id"], "quote": ctx["span_text"]}],
                "rationale": f"Assessment for {cid.value}",
                "missing_information": [],
            }
            for cid in CriterionId
        ]
        # Discriminatory reason note mentioning age/tuổi
        change_reasons = {
            "python_backend": {
                "reason_code": "other",
                "note": "Ứng viên này còn quá trẻ, tuổi đời ít nên chưa phù hợp vai trò chính.",
            }
        }

        res = await client.post(
            f"/api/v1/applications/{ctx['app_id']}/hr-revisions",
            json={
                "expected_application_version": 1,
                "source_snapshot_ref": {
                    "document_id": str(ctx["doc_id"]),
                    "sanitized_version_id": str(ctx["sanitized_id"]),
                    "rubric_version_id": str(ctx["rubric_id"]),
                    "application_generation": 1,
                },
                "criteria": criteria,
                "change_reasons": change_reasons,
                "summary_reason": "Xem xét lại hồ sơ",
            },
        )
        assert res.status_code == 422
        assert "FORBIDDEN_DEMOGRAPHIC_ATTRIBUTE" in res.json()["detail"]


@pytest.mark.asyncio
async def test_finalize_hr_revision_immutability(test_session_factory, sample_docx_cv):
    """Finalized revision becomes immutable with content hash."""
    async with test_session_factory() as session:
        ctx = await setup_test_context(session, sample_docx_cv)

    transport = ASGITransport(app=app)
    cookies = {SESSION_COOKIE_NAME: ctx["r_token"]}
    headers = {"X-CSRF-Token": ctx["r_csrf"]}

    async with AsyncClient(transport=transport, base_url="http://test", cookies=cookies, headers=headers) as client:
        criteria = [
            {
                "criterion_id": cid.value,
                "status": "assessed",
                "score": 3,
                "evidence": [{"span_id": ctx["span_id"], "quote": ctx["span_text"]}],
                "rationale": f"Assessment for {cid.value}",
                "missing_information": [],
            }
            for cid in CriterionId
        ]
        res = await client.post(
            f"/api/v1/applications/{ctx['app_id']}/hr-revisions",
            json={
                "expected_application_version": 1,
                "source_snapshot_ref": {
                    "document_id": str(ctx["doc_id"]),
                    "sanitized_version_id": str(ctx["sanitized_id"]),
                    "rubric_version_id": str(ctx["rubric_id"]),
                    "application_generation": 1,
                },
                "criteria": criteria,
                "change_reasons": {},
                "summary_reason": "Complete independent review.",
            },
        )
        rev_id = res.json()["id"]

        # Finalize
        fin_res = await client.post(
            f"/api/v1/hr-revisions/{rev_id}/finalize",
            json={
                "expected_application_version": 1,
                "expected_rubric_version_id": str(ctx["rubric_id"]),
            },
        )
        assert fin_res.status_code == 200
        fin_data = fin_res.json()
        assert fin_data["status"] == "finalized"
        assert fin_data["content_hash"] is not None
        assert len(fin_data["content_hash"]) == 64

        # Attempt to edit finalized revision -> 409 Conflict
        edit_res = await client.put(
            f"/api/v1/hr-revisions/{rev_id}",
            json={
                "expected_revision_version": 1,
                "criteria": criteria,
                "change_reasons": {},
                "summary_reason": "Trying to edit finalized",
            },
        )
        assert edit_res.status_code == 409
        assert "REVISION_FINALIZED" in edit_res.json()["detail"]


@pytest.mark.asyncio
async def test_reviewer_cannot_make_hiring_decision_but_owner_can(test_session_factory, sample_docx_cv):
    """Reviewer cannot make hiring decision (403), Owner can attest and decide (201)."""
    async with test_session_factory() as session:
        ctx = await setup_test_context(session, sample_docx_cv)

    transport = ASGITransport(app=app)
    r_client = AsyncClient(transport=transport, base_url="http://test", cookies={SESSION_COOKIE_NAME: ctx["r_token"]}, headers={"X-CSRF-Token": ctx["r_csrf"]})
    o_client = AsyncClient(transport=transport, base_url="http://test", cookies={SESSION_COOKIE_NAME: ctx["o_token"]}, headers={"X-CSRF-Token": ctx["o_csrf"]})

    async with r_client, o_client:
        # 1. Create and finalize HR revision
        criteria = [
            {
                "criterion_id": cid.value,
                "status": "assessed",
                "score": 3,
                "evidence": [{"span_id": ctx["span_id"], "quote": ctx["span_text"]}],
                "rationale": f"Assessment for {cid.value}",
                "missing_information": [],
            }
            for cid in CriterionId
        ]
        rev_res = await r_client.post(
            f"/api/v1/applications/{ctx['app_id']}/hr-revisions",
            json={
                "expected_application_version": 1,
                "source_snapshot_ref": {
                    "document_id": str(ctx["doc_id"]),
                    "sanitized_version_id": str(ctx["sanitized_id"]),
                    "rubric_version_id": str(ctx["rubric_id"]),
                    "application_generation": 1,
                },
                "criteria": criteria,
                "change_reasons": {},
                "summary_reason": "Reviewer finalized evaluation.",
            },
        )
        rev_id = rev_res.json()["id"]
        await r_client.post(
            f"/api/v1/hr-revisions/{rev_id}/finalize",
            json={
                "expected_application_version": 1,
                "expected_rubric_version_id": str(ctx["rubric_id"]),
            },
        )

        # 2. Reviewer signs attestation
        r_att_res = await r_client.post(
            f"/api/v1/applications/{ctx['app_id']}/review-attestations",
            json={
                "decision_basis": "assessment_review",
                "effective_result": {"kind": "hr_revision", "id": rev_id},
                "reviewed_criterion_ids": [cid.value for cid in CriterionId],
                "acknowledged": True,
            },
        )
        assert r_att_res.status_code == 201
        r_att_id = r_att_res.json()["id"]

        # 3. Reviewer tries to create hiring decision -> 403 Forbidden!
        r_dec_res = await r_client.post(
            f"/api/v1/applications/{ctx['app_id']}/decisions",
            json={
                "decision_basis": "assessment_review",
                "outcome": "advance",
                "reason": "Mời ứng viên vào vòng phỏng vấn kỹ thuật trực tiếp.",
                "attestation_id": r_att_id,
            },
        )
        assert r_dec_res.status_code == 403
        assert "Chỉ Owner" in r_dec_res.json()["detail"]

        # 4. Owner signs own attestation
        o_att_res = await o_client.post(
            f"/api/v1/applications/{ctx['app_id']}/review-attestations",
            json={
                "decision_basis": "assessment_review",
                "effective_result": {"kind": "hr_revision", "id": rev_id},
                "reviewed_criterion_ids": [cid.value for cid in CriterionId],
                "acknowledged": True,
            },
        )
        assert o_att_res.status_code == 201
        o_att_id = o_att_res.json()["id"]

        # 5. Owner makes hiring decision -> 201 Created!
        o_dec_res = await o_client.post(
            f"/api/v1/applications/{ctx['app_id']}/decisions",
            json={
                "decision_basis": "assessment_review",
                "outcome": "advance",
                "reason": "Đã xem xét bản đánh giá HR và đồng ý mời ứng viên vào vòng phỏng vấn.",
                "attestation_id": o_att_id,
                "expected_previous_decision_id": None,
                "expected_rubric_version_id": str(ctx["rubric_id"]),
            },
        )
        assert o_dec_res.status_code == 201
        dec_data = o_dec_res.json()
        assert dec_data["sequence_no"] == 1
        assert dec_data["outcome"] == "advance"
        assert dec_data["decision_basis"] == "assessment_review"

        # 6. Verify application current_decision_id updated
        async with test_session_factory() as session:
            stmt_app = select(Application).where(Application.id == ctx["app_id"])
            updated_app = (await session.execute(stmt_app)).scalar_one()
            assert str(updated_app.current_decision_id) == dec_data["id"]
            assert updated_app.row_version == 2


@pytest.mark.asyncio
async def test_technical_information_request_basis(test_session_factory, sample_docx_cv):
    """Technical error allows only request_information outcome."""
    async with test_session_factory() as session:
        ctx = await setup_test_context(session, sample_docx_cv)

    transport = ASGITransport(app=app)
    cookies = {SESSION_COOKIE_NAME: ctx["o_token"]}
    headers = {"X-CSRF-Token": ctx["o_csrf"]}

    async with AsyncClient(transport=transport, base_url="http://test", cookies=cookies, headers=headers) as client:
        # 1. Attest technical issue
        att_res = await client.post(
            f"/api/v1/applications/{ctx['app_id']}/review-attestations",
            json={
                "decision_basis": "technical_information_request",
                "document_id": str(ctx["doc_id"]),
                "expected_document_sha256": "c" * 64,
                "failure_ref": {"kind": "document", "id": str(ctx["doc_id"])},
                "technical_failure_code": "OCR_QUALITY_FAILED",
                "reviewed_criterion_ids": [],
                "acknowledged": True,
            },
        )
        assert att_res.status_code == 201
        att_id = att_res.json()["id"]

        # 2. Attempt advance outcome -> 422 Invalid Outcome
        bad_dec = await client.post(
            f"/api/v1/applications/{ctx['app_id']}/decisions",
            json={
                "decision_basis": "technical_information_request",
                "outcome": "advance",
                "reason": "File bị mờ nhưng vẫn muốn mời phỏng vấn để kiểm tra.",
                "attestation_id": att_id,
            },
        )
        assert bad_dec.status_code == 422
        assert "INVALID_OUTCOME_FOR_TECHNICAL_REQUEST" in bad_dec.json()["detail"]

        # 3. Valid request_information outcome -> 201 Created
        good_dec = await client.post(
            f"/api/v1/applications/{ctx['app_id']}/decisions",
            json={
                "decision_basis": "technical_information_request",
                "outcome": "request_information",
                "reason": "File CV bị lỗi OCR chất lượng kém, yêu cầu ứng viên tải lại bản scan rõ nét hơn.",
                "attestation_id": att_id,
            },
        )
        assert good_dec.status_code == 201
        assert good_dec.json()["outcome"] == "request_information"


@pytest.mark.asyncio
async def test_manual_document_review_with_raw_grant(test_session_factory, sample_docx_cv):
    """Manual document review requires active raw_cv grant and override reason."""
    async with test_session_factory() as session:
        ctx = await setup_test_context(session, sample_docx_cv)

    transport = ASGITransport(app=app)
    cookies = {SESSION_COOKIE_NAME: ctx["o_token"]}
    headers = {"X-CSRF-Token": ctx["o_csrf"]}

    async with AsyncClient(transport=transport, base_url="http://test", cookies=cookies, headers=headers) as client:
        # 1. Without grant -> 403 Forbidden
        att_no_grant = await client.post(
            f"/api/v1/applications/{ctx['app_id']}/review-attestations",
            json={
                "decision_basis": "manual_document_review",
                "document_id": str(ctx["doc_id"]),
                "expected_document_sha256": "c" * 64,
                "expected_rubric_version_id": str(ctx["rubric_id"]),
                "reviewed_criterion_ids": [cid.value for cid in CriterionId],
                "manual_evidence_refs": [
                    {
                        "criterion_id": "python_backend",
                        "document_id": str(ctx["doc_id"]),
                        "page_number": 1,
                        "section_label": "Kinh nghiệm",
                        "note": "Đã xem tài liệu gốc xác nhận năng lực.",
                    }
                ],
                "acknowledged": True,
            },
        )
        assert att_no_grant.status_code == 403

        # 2. Grant raw access
        grant_res = await client.post(
            f"/api/v1/applications/{ctx['app_id']}/raw-grants",
            json={
                "grantee_user_id": str(ctx["owner"].id),
                "scopes": ["raw_cv"],
                "expires_at": (datetime.now(timezone.utc) + timedelta(hours=2)).isoformat(),
                "reason": "Owner manual CV review",
            },
        )
        assert grant_res.status_code == 201

        # 3. Now manual attestation succeeds
        att_res = await client.post(
            f"/api/v1/applications/{ctx['app_id']}/review-attestations",
            json={
                "decision_basis": "manual_document_review",
                "document_id": str(ctx["doc_id"]),
                "expected_document_sha256": "c" * 64,
                "expected_rubric_version_id": str(ctx["rubric_id"]),
                "reviewed_criterion_ids": [cid.value for cid in CriterionId],
                "manual_evidence_refs": [
                    {
                        "criterion_id": "python_backend",
                        "document_id": str(ctx["doc_id"]),
                        "page_number": 1,
                        "section_label": "Kinh nghiệm",
                        "note": "Đã xem tài liệu gốc xác nhận năng lực.",
                    }
                ],
                "acknowledged": True,
            },
        )
        assert att_res.status_code == 201
        att_id = att_res.json()["id"]

        # 4. Decision without override_reason -> 422
        bad_dec = await client.post(
            f"/api/v1/applications/{ctx['app_id']}/decisions",
            json={
                "decision_basis": "manual_document_review",
                "outcome": "advance",
                "reason": "Đã đọc hồ sơ và quyết định phỏng vấn trực tiếp ứng viên.",
                "attestation_id": att_id,
            },
        )
        assert bad_dec.status_code == 422
        assert "OVERRIDE_REASON_REQUIRED" in bad_dec.json()["detail"]

        # 5. Decision with override_reason -> 201
        good_dec = await client.post(
            f"/api/v1/applications/{ctx['app_id']}/decisions",
            json={
                "decision_basis": "manual_document_review",
                "outcome": "advance",
                "reason": "Đã đọc trực tiếp file tài liệu gốc và đối chiếu cả 6 tiêu chí.",
                "override_reason": "Quyết định đánh giá thủ công hoàn toàn trên file gốc theo quy trình bypass.",
                "attestation_id": att_id,
            },
        )
        assert good_dec.status_code == 201
        assert good_dec.json()["decision_basis"] == "manual_document_review"


@pytest.mark.asyncio
async def test_diverging_assessment_decision_requires_override_reason(test_session_factory, sample_docx_cv):
    """When recommendation is consider_next_round, choosing not_advance requires override_reason."""
    async with test_session_factory() as session:
        ctx = await setup_test_context(session, sample_docx_cv)

    transport = ASGITransport(app=app)
    cookies = {SESSION_COOKIE_NAME: ctx["o_token"]}
    headers = {"X-CSRF-Token": ctx["o_csrf"]}

    async with AsyncClient(transport=transport, base_url="http://test", cookies=cookies, headers=headers) as client:
        # Create and finalize HR revision with score 3 (recommendation: consider_next_round)
        criteria = [
            {
                "criterion_id": cid.value,
                "status": "assessed",
                "score": 3,
                "evidence": [{"span_id": ctx["span_id"], "quote": ctx["span_text"]}],
                "rationale": f"Assessment for {cid.value}",
                "missing_information": [],
            }
            for cid in CriterionId
        ]
        rev_res = await client.post(
            f"/api/v1/applications/{ctx['app_id']}/hr-revisions",
            json={
                "expected_application_version": 1,
                "source_snapshot_ref": {
                    "document_id": str(ctx["doc_id"]),
                    "sanitized_version_id": str(ctx["sanitized_id"]),
                    "rubric_version_id": str(ctx["rubric_id"]),
                    "application_generation": 1,
                },
                "criteria": criteria,
                "change_reasons": {},
                "summary_reason": "High score evaluation.",
            },
        )
        rev_id = rev_res.json()["id"]
        await client.post(
            f"/api/v1/hr-revisions/{rev_id}/finalize",
            json={
                "expected_application_version": 1,
                "expected_rubric_version_id": str(ctx["rubric_id"]),
            },
        )

        # Attest
        att_res = await client.post(
            f"/api/v1/applications/{ctx['app_id']}/review-attestations",
            json={
                "decision_basis": "assessment_review",
                "effective_result": {"kind": "hr_revision", "id": rev_id},
                "reviewed_criterion_ids": [cid.value for cid in CriterionId],
                "acknowledged": True,
            },
        )
        att_id = att_res.json()["id"]

        # Attempt not_advance without override_reason -> 422
        bad_dec = await client.post(
            f"/api/v1/applications/{ctx['app_id']}/decisions",
            json={
                "decision_basis": "assessment_review",
                "outcome": "not_advance",
                "reason": "Dù điểm cao nhưng ứng viên yêu cầu mức lương vượt khung ngân sách.",
                "attestation_id": att_id,
            },
        )
        assert bad_dec.status_code == 422
        assert "OVERRIDE_REASON_REQUIRED" in bad_dec.json()["detail"]

        # With override_reason -> 201
        good_dec = await client.post(
            f"/api/v1/applications/{ctx['app_id']}/decisions",
            json={
                "decision_basis": "assessment_review",
                "outcome": "not_advance",
                "reason": "Dù điểm chuyên môn đạt yêu cầu nhưng không thể sắp xếp lịch phỏng vấn phù hợp.",
                "override_reason": "Giải trình từ chối ứng viên có điểm khuyến nghị đạt do định biên nhân sự đã đủ.",
                "attestation_id": att_id,
                "expected_previous_decision_id": None,
                "expected_rubric_version_id": str(ctx["rubric_id"]),
            },
        )
        assert good_dec.status_code == 201
        assert good_dec.json()["outcome"] == "not_advance"


@pytest.mark.asyncio
async def test_attestation_mismatch_and_concurrency_conflict(test_session_factory, sample_docx_cv):
    """Owner cannot borrow Reviewer's attestation, and outdated expected_previous_decision triggers 409."""
    async with test_session_factory() as session:
        ctx = await setup_test_context(session, sample_docx_cv)

    transport = ASGITransport(app=app)
    r_client = AsyncClient(transport=transport, base_url="http://test", cookies={SESSION_COOKIE_NAME: ctx["r_token"]}, headers={"X-CSRF-Token": ctx["r_csrf"]})
    o_client = AsyncClient(transport=transport, base_url="http://test", cookies={SESSION_COOKIE_NAME: ctx["o_token"]}, headers={"X-CSRF-Token": ctx["o_csrf"]})

    async with r_client, o_client:
        criteria = [
            {
                "criterion_id": cid.value,
                "status": "assessed",
                "score": 3,
                "evidence": [{"span_id": ctx["span_id"], "quote": ctx["span_text"]}],
                "rationale": f"Assessment for {cid.value}",
                "missing_information": [],
            }
            for cid in CriterionId
        ]
        rev_res = await r_client.post(
            f"/api/v1/applications/{ctx['app_id']}/hr-revisions",
            json={
                "expected_application_version": 1,
                "source_snapshot_ref": {
                    "document_id": str(ctx["doc_id"]),
                    "sanitized_version_id": str(ctx["sanitized_id"]),
                    "rubric_version_id": str(ctx["rubric_id"]),
                    "application_generation": 1,
                },
                "criteria": criteria,
                "change_reasons": {},
                "summary_reason": "Base evaluation.",
            },
        )
        rev_id = rev_res.json()["id"]
        await r_client.post(
            f"/api/v1/hr-revisions/{rev_id}/finalize",
            json={
                "expected_application_version": 1,
                "expected_rubric_version_id": str(ctx["rubric_id"]),
            },
        )

        # Reviewer attestation
        r_att = await r_client.post(
            f"/api/v1/applications/{ctx['app_id']}/review-attestations",
            json={
                "decision_basis": "assessment_review",
                "effective_result": {"kind": "hr_revision", "id": rev_id},
                "reviewed_criterion_ids": [cid.value for cid in CriterionId],
                "acknowledged": True,
            },
        )
        r_att_id = r_att.json()["id"]

        # Owner tries to use Reviewer's attestation -> 403 ATTESTATION_ACTOR_MISMATCH
        mismatch_res = await o_client.post(
            f"/api/v1/applications/{ctx['app_id']}/decisions",
            json={
                "decision_basis": "assessment_review",
                "outcome": "advance",
                "reason": "Chấp thuận đánh giá của reviewer.",
                "attestation_id": r_att_id,
            },
        )
        assert mismatch_res.status_code == 403
        assert "ATTESTATION_ACTOR_MISMATCH" in mismatch_res.json()["detail"]

        # Owner signs own attestation
        o_att = await o_client.post(
            f"/api/v1/applications/{ctx['app_id']}/review-attestations",
            json={
                "decision_basis": "assessment_review",
                "effective_result": {"kind": "hr_revision", "id": rev_id},
                "reviewed_criterion_ids": [cid.value for cid in CriterionId],
                "acknowledged": True,
            },
        )
        o_att_id = o_att.json()["id"]

        # First decision succeeds
        d1 = await o_client.post(
            f"/api/v1/applications/{ctx['app_id']}/decisions",
            json={
                "decision_basis": "assessment_review",
                "outcome": "advance",
                "reason": "Quyết định phỏng vấn sau khi tự rà soát hồ sơ cẩn thận.",
                "attestation_id": o_att_id,
                "expected_previous_decision_id": None,
            },
        )
        assert d1.status_code == 201

        # Second decision trying to say expected_previous_decision_id is None -> 409 DECISION_CONFLICT
        d2 = await o_client.post(
            f"/api/v1/applications/{ctx['app_id']}/decisions",
            json={
                "decision_basis": "assessment_review",
                "outcome": "not_advance",
                "reason": "Thay đổi quyết định do thay đổi kế hoạch.",
                "override_reason": "Giải trình thay đổi quyết định đã có từ trước.",
                "attestation_id": o_att_id,
                "expected_previous_decision_id": None,
            },
        )
        assert d2.status_code == 409
        assert "DECISION_CONFLICT" in d2.json()["detail"]
