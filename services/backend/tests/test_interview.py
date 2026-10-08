"""Comprehensive tests for Task B14: Rubric Agent, Interview Question Banks, and Interview Agent."""
from __future__ import annotations

from datetime import datetime, timezone
import io
import json
import uuid
import docx
import pytest
from httpx import ASGITransport, AsyncClient
from sqlalchemy import delete, select

from app.db.models import (
    Application,
    AssessmentRun,
    BudgetReservation,
    Candidate,
    CriterionAssessment,
    CriterionEvidence,
    Decision,
    Document,
    HRRevision,
    InterviewDraft,
    InterviewQuestionBank,
    InterviewRevision,
    InterviewScorecard,
    JDVersion,
    Job,
    Requisition,
    RequisitionMembership,
    ReviewAttestation,
    RubricCriterion,
    RubricVersion,
    SanitizedVersion,
    SourceSpan,
    User,
    UserAccountRole,
)
from app.domain.authorization import SESSION_COOKIE_NAME
from app.domain.enums import (
    AccountRole,
    CriterionId,
    CriterionOutcome,
    DocumentSafetyStatus,
    JobStatus,
    JobType,
    MembershipRole,
    Recommendation,
    RequisitionStatus,
    RubricStatus,
    SanitizedVersionStatus,
    UserStatus,
)
from app.domain.security import hash_password
from app.main import app
from app.services.auth import create_session
from app.services.interview import execute_interview_job
from app.services.llm.provider import MockLLMProvider
from app.services.worker import claim_next_job


@pytest.fixture(autouse=True)
async def cleanup_interview_tables(test_session_factory):
    async with test_session_factory() as session:
        await session.execute(delete(InterviewRevision))
        await session.execute(delete(InterviewDraft))
        await session.execute(delete(InterviewQuestionBank))
        await session.execute(delete(Decision))
        await session.execute(delete(ReviewAttestation))
        await session.execute(delete(HRRevision))
        await session.execute(delete(CriterionEvidence))
        await session.execute(delete(CriterionAssessment))
        await session.execute(delete(AssessmentRun))
        await session.execute(delete(BudgetReservation))
        await session.execute(delete(Job))
        await session.execute(delete(SourceSpan))
        await session.commit()
    yield
    async with test_session_factory() as session:
        await session.execute(delete(InterviewRevision))
        await session.execute(delete(InterviewDraft))
        await session.execute(delete(InterviewQuestionBank))
        await session.execute(delete(Decision))
        await session.execute(delete(ReviewAttestation))
        await session.execute(delete(HRRevision))
        await session.execute(delete(CriterionEvidence))
        await session.execute(delete(CriterionAssessment))
        await session.execute(delete(AssessmentRun))
        await session.execute(delete(BudgetReservation))
        await session.execute(delete(Job))
        await session.execute(delete(SourceSpan))
        await session.commit()


@pytest.fixture
def sample_docx_cv() -> bytes:
    doc = docx.Document()
    doc.add_paragraph("Kinh nghiệm làm việc chuyên sâu Backend Python và FastAPI.")
    doc.add_paragraph("Thiết kế kiến trúc hệ thống microservices chịu tải cao và tối ưu cơ sở dữ liệu PostgreSQL.")
    doc.add_paragraph("Triển khai hạ tầng container với Docker và Kubernetes.")
    doc.add_paragraph("Kỹ năng giao tiếp và làm việc nhóm hiệu quả.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def setup_interview_test_context(session, sample_docx_cv):
    from app.services.requisition import get_or_create_default_org
    org = await get_or_create_default_org(session)

    owner = User(
        id=uuid.uuid4(),
        login_name=f"owner_{uuid.uuid4().hex[:8]}",
        display_name="Owner HR",
        password_hash=hash_password("Password123!"),
        status=UserStatus.ACTIVE,
    )
    reviewer = User(
        id=uuid.uuid4(),
        login_name=f"rev_{uuid.uuid4().hex[:8]}",
        display_name="Reviewer HR",
        password_hash=hash_password("Password123!"),
        status=UserStatus.ACTIVE,
    )
    session.add_all([owner, reviewer])
    await session.flush()
    session.add(UserAccountRole(user_id=owner.id, role=AccountRole.RECRUITER))
    session.add(UserAccountRole(user_id=reviewer.id, role=AccountRole.RECRUITER))

    o_token, o_csrf, _ = await create_session(session, owner)
    r_token, r_csrf, _ = await create_session(session, reviewer)

    req = Requisition(
        id=uuid.uuid4(),
        organization_id=org.id,
        title="Senior Python Engineer",
        status=RequisitionStatus.OPEN,
        row_version=1,
    )
    session.add(req)
    await session.flush()
    session.add(RequisitionMembership(requisition_id=req.id, user_id=owner.id, membership_role=MembershipRole.OWNER))
    session.add(RequisitionMembership(requisition_id=req.id, user_id=reviewer.id, membership_role=MembershipRole.REVIEWER))
    await session.flush()

    jd = JDVersion(
        id=uuid.uuid4(),
        requisition_id=req.id,
        version_no=1,
        source_text="Senior Python Backend Engineer JD",
        text_hash="a" * 64,
        created_by=owner.id,
    )
    session.add(jd)
    await session.flush()
    req.current_jd_version_id = jd.id

    rubric = RubricVersion(
        id=uuid.uuid4(),
        requisition_id=req.id,
        jd_version_id=jd.id,
        version_no=1,
        status=RubricStatus.APPROVED,
        content_hash="b" * 64,
        approved_by=owner.id,
        approved_at=datetime.now(timezone.utc),
    )
    session.add(rubric)
    await session.flush()
    req.current_rubric_version_id = rubric.id

    for cid in CriterionId:
        crit = RubricCriterion(
            rubric_version_id=rubric.id,
            criterion_id=cid.value,
            label_vi=cid.name,
            weight=20 if cid == CriterionId.PYTHON_BACKEND else 16,
            description_vi=f"Description for {cid.name}",
            anchors={"0": "None", "1": "Basic", "2": "Intermediate", "3": "Advanced", "4": "Expert"},
        )
        session.add(crit)

    cand = Candidate(id=uuid.uuid4(), organization_id=org.id, public_label=f"CAND-{uuid.uuid4().hex[:6].upper()}", status="active")
    session.add(cand)
    await session.flush()

    app_obj = Application(
        id=uuid.uuid4(),
        requisition_id=req.id,
        candidate_id=cand.id,
        status="active",
        generation=1,
        row_version=1,
    )
    session.add(app_obj)
    await session.flush()

    doc = Document(
        id=uuid.uuid4(),
        application_id=app_obj.id,
        version_no=1,
        kind="cv",
        original_name_private="cv.docx",
        mime_verified="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        byte_size=len(sample_docx_cv),
        sha256="c" * 64,
        blob_key="docs/test.docx",
        ingestion_status="parsed",
        safety_status=DocumentSafetyStatus.PASSED,
        created_by=owner.id,
    )
    session.add(doc)
    await session.flush()
    app_obj.current_document_id = doc.id

    sanitized = SanitizedVersion(
        id=uuid.uuid4(),
        application_id=app_obj.id,
        document_id=doc.id,
        version_no=1,
        status=SanitizedVersionStatus.APPROVED,
        canonical_text="Kinh nghiệm làm việc chuyên sâu Backend Python và FastAPI.\nThiết kế microservices PostgreSQL.",
        sha256="d" * 64,
        approved_by=owner.id,
        approved_at=datetime.now(timezone.utc),
    )
    session.add(sanitized)
    await session.flush()
    app_obj.current_sanitized_version_id = sanitized.id

    span_id = f"spn_{uuid.uuid4().hex[:24]}"
    span1 = SourceSpan(
        span_id=span_id,
        full_hash="e" * 64,
        sanitized_version_id=sanitized.id,
        start_cp=0,
        end_cp=58,
        page_number=1,
        language="vi",
        text="Kinh nghiệm làm việc chuyên sâu Backend Python và FastAPI.",
    )
    session.add(span1)

    # Completed Assessment Run for interview draft input
    run_job = Job(
        id=uuid.uuid4(),
        type=JobType.ASSESS_APPLICATION,
        target_type="application",
        status=JobStatus.SUCCEEDED,
        target_id=app_obj.id,
        input_snapshot_hash="e3b0c44298fc1c149afbf4c8996fb92427ae41e4649b934ca495991b7852b855",
        payload_ref={},
    )
    session.add(run_job)
    await session.flush()

    run = AssessmentRun(
        id=uuid.uuid4(),
        application_id=app_obj.id,
        job_id=run_job.id,
        run_no=1,
        status="succeeded",
        snapshot={},
        snapshot_hash="f" * 64,
        application_generation=1,
        document_id=doc.id,
        sanitized_version_id=sanitized.id,
        rubric_version_id=rubric.id,
        observed_score=75.0,
        coverage=1.0,
        comparable_score=75.0,
        recommendation=Recommendation.CONSIDER_NEXT_ROUND,
    )
    session.add(run)
    await session.flush()
    app_obj.current_assessment_run_id = run.id

    for cid in CriterionId:
        ca = CriterionAssessment(
            run_id=run.id,
            criterion_id=cid.value,
            status=CriterionOutcome.ASSESSED,
            score=3,
            rationale=f"Evaluated competency at level 3 in {cid.value}",
            missing_information=[],
        )
        session.add(ca)

    await session.commit()

    return {
        "owner": owner,
        "o_token": o_token,
        "o_csrf": o_csrf,
        "reviewer": reviewer,
        "r_token": r_token,
        "r_csrf": r_csrf,
        "req_id": req.id,
        "rubric_id": rubric.id,
        "app_id": app_obj.id,
        "doc_id": doc.id,
        "sanitized_id": sanitized.id,
        "span_id": span_id,
        "span_text": span1.text,
        "run_id": run.id,
    }


# ---------------- Question Bank Tests ---------------- #


@pytest.mark.asyncio
async def test_question_bank_seed_import_and_rbac(test_session_factory, sample_docx_cv):
    """Import seed question bank creates DRAFT (not approved), Reviewer cannot approve, Owner can approve."""
    async with test_session_factory() as session:
        ctx = await setup_interview_test_context(session, sample_docx_cv)

    transport = ASGITransport(app=app)
    r_client = AsyncClient(transport=transport, base_url="http://test", cookies={SESSION_COOKIE_NAME: ctx["r_token"]}, headers={"X-CSRF-Token": ctx["r_csrf"]})
    o_client = AsyncClient(transport=transport, base_url="http://test", cookies={SESSION_COOKIE_NAME: ctx["o_token"]}, headers={"X-CSRF-Token": ctx["o_csrf"]})

    async with r_client, o_client:
        # 1. Reviewer tries to create question bank -> 403 Forbidden
        bad_create = await r_client.post(
            f"/api/v1/rubrics/{ctx['rubric_id']}/interview-question-banks",
            json={"source": "seed"},
        )
        assert bad_create.status_code == 403

        # 2. Owner imports seed bank -> 201 Created in DRAFT status
        seed_res = await o_client.post(
            f"/api/v1/rubrics/{ctx['rubric_id']}/interview-question-banks",
            json={"source": "seed"},
        )
        assert seed_res.status_code == 201
        data = seed_res.json()
        assert data["status"] == "draft"
        assert len(data["questions"]) == 6
        assert data["version_no"] == 1
        bank_id = data["id"]

        # 3. Reviewer tries to approve bank -> 403 Forbidden
        bad_appr = await r_client.post(
            f"/api/v1/interview-question-banks/{bank_id}/approve",
            json={"expected_rubric_version_id": str(ctx["rubric_id"]), "acknowledged": True},
        )
        assert bad_appr.status_code == 403

        # 4. Owner approves bank -> 200 OK
        appr_res = await o_client.post(
            f"/api/v1/interview-question-banks/{bank_id}/approve",
            json={"expected_rubric_version_id": str(ctx["rubric_id"]), "acknowledged": True},
        )
        assert appr_res.status_code == 200
        assert appr_res.json()["status"] == "approved"
        assert appr_res.json()["approved_by"] is not None


@pytest.mark.asyncio
async def test_question_bank_anti_discrimination_scan(test_session_factory, sample_docx_cv):
    """Editing question bank with forbidden demographic questions is rejected with 422."""
    async with test_session_factory() as session:
        ctx = await setup_interview_test_context(session, sample_docx_cv)

    transport = ASGITransport(app=app)
    cookies = {SESSION_COOKIE_NAME: ctx["o_token"]}
    headers = {"X-CSRF-Token": ctx["o_csrf"]}

    async with AsyncClient(transport=transport, base_url="http://test", cookies=cookies, headers=headers) as client:
        # Import draft bank
        res = await client.post(
            f"/api/v1/rubrics/{ctx['rubric_id']}/interview-question-banks",
            json={"source": "seed"},
        )
        bank_id = res.json()["id"]
        questions = res.json()["questions"]

        # Inject age question
        questions[0]["question_vi"] = "Bạn sinh năm bao nhiêu và tuổi của bạn có ảnh hưởng tới công việc không?"

        up_res = await client.put(
            f"/api/v1/interview-question-banks/{bank_id}",
            json={"questions": questions, "change_reason": "Thêm câu hỏi rà soát thông tin cá nhân"},
        )
        assert up_res.status_code == 422
        assert "FORBIDDEN_DEMOGRAPHIC_ATTRIBUTE" in up_res.json()["detail"]


# ---------------- Interview Draft & Follow-ups Tests ---------------- #


@pytest.mark.asyncio
async def test_interview_draft_generation_and_worker_execution(test_session_factory, sample_docx_cv):
    """Enqueue interview draft, worker runs MockLLMProvider, generates followups conforming to schema."""
    async with test_session_factory() as session:
        ctx = await setup_interview_test_context(session, sample_docx_cv)

    transport = ASGITransport(app=app)
    cookies = {SESSION_COOKIE_NAME: ctx["o_token"]}
    headers = {"X-CSRF-Token": ctx["o_csrf"]}

    async with AsyncClient(transport=transport, base_url="http://test", cookies=cookies, headers=headers) as client:
        # 1. Setup approved question bank
        bank_res = await client.post(
            f"/api/v1/rubrics/{ctx['rubric_id']}/interview-question-banks",
            json={"source": "seed"},
        )
        bank_id = bank_res.json()["id"]
        await client.post(
            f"/api/v1/interview-question-banks/{bank_id}/approve",
            json={"expected_rubric_version_id": str(ctx["rubric_id"]), "acknowledged": True},
        )

        # 2. Enqueue interview draft
        draft_res = await client.post(
            f"/api/v1/applications/{ctx['app_id']}/interview-drafts",
            json={
                "effective_result": {"kind": "assessment_run", "id": str(ctx["run_id"])},
                "expected_question_bank_id": bank_id,
            },
        )
        assert draft_res.status_code == 202
        draft_id = draft_res.json()["id"]
        job_id = draft_res.json()["job_id"]

    # 3. Worker executes with mock provider returning valid followups
    mock_payload = {
        "followups": [
            {
                "criterion_id": "python_backend",
                "question_vi": "Bạn có thể giải thích chi tiết hơn về cách tổ chức mô-đun trong ứng dụng FastAPI?",
                "purpose_vi": "Làm rõ cấu trúc mã nguồn và tính phân tách trách nhiệm.",
                "source_span_ids": [ctx["span_id"]],
                "answer_indicators": ["Giải thích mô hình router", "Nêu xử lý exception handler"],
            }
        ]
    }
    mock_llm = MockLLMProvider(custom_content=json.dumps(mock_payload))

    async with test_session_factory() as session:
        await execute_interview_job(session, job_id=uuid.UUID(job_id), provider_override=mock_llm)
        await session.commit()

    # 4. Fetch draft detail
    async with AsyncClient(transport=transport, base_url="http://test", cookies=cookies, headers=headers) as client:
        detail_res = await client.get(f"/api/v1/interview-drafts/{draft_id}")
        assert detail_res.status_code == 200
        detail = detail_res.json()
        assert detail["status"] == "succeeded"
        assert len(detail["core_questions"]) == 6  # Read-only core questions
        assert len(detail["ai_followups"]) == 1
        assert detail["ai_followups"][0]["criterion_id"] == "python_backend"
        assert detail["is_stale"] is False

        # 5. HR saves candidate-specific follow-up revisions
        edited_followups = [
            {
                "criterion_id": "python_backend",
                "question_vi": "Bạn xử lý migration cơ sở dữ liệu với Alembic trong FastAPI như thế nào?",
                "purpose_vi": "Kiểm tra kinh nghiệm quản lý database schema.",
                "source_span_ids": [ctx["span_id"]],
                "answer_indicators": ["Nêu quy trình autogenerate", "Kiểm tra downgrade script"],
            }
        ]
        rev_res = await client.post(
            f"/api/v1/interview-drafts/{draft_id}/revisions",
            json={
                "expected_previous_revision_id": None,
                "followups": edited_followups,
                "change_reason": "HR bổ sung câu hỏi về database migration cho phù hợp thực tế dự án.",
            },
        )
        assert rev_res.status_code == 201
        assert rev_res.json()["revision_no"] == 1

        # 6. Reload draft to verify HR revision is returned as latest
        reload_res = await client.get(f"/api/v1/interview-drafts/{draft_id}")
        assert reload_res.status_code == 200
        latest_rev = reload_res.json()["latest_revision"]
        assert latest_rev is not None
        assert latest_rev["followups"][0]["question_vi"] == edited_followups[0]["question_vi"]
        # Verify core questions STILL unaltered
        assert len(reload_res.json()["core_questions"]) == 6


@pytest.mark.asyncio
async def test_interview_followups_do_not_require_a_question_bank(test_session_factory, sample_docx_cv):
    async with test_session_factory() as session:
        ctx = await setup_interview_test_context(session, sample_docx_cv)

    transport = ASGITransport(app=app)
    cookies = {SESSION_COOKIE_NAME: ctx["o_token"]}
    headers = {"X-CSRF-Token": ctx["o_csrf"]}
    async with AsyncClient(transport=transport, base_url="http://test", cookies=cookies, headers=headers) as client:
        draft_res = await client.post(
            f"/api/v1/applications/{ctx['app_id']}/interview-drafts",
            json={"effective_result": {"kind": "assessment_run", "id": str(ctx["run_id"])}},
        )
        assert draft_res.status_code == 202
        draft_id = draft_res.json()["id"]
        job_id = draft_res.json()["job_id"]
        assert draft_res.json()["question_bank_id"] is None
        assert draft_res.json()["core_questions"] == []

    mock_llm = MockLLMProvider(custom_content=json.dumps({
        "followups": [{
            "criterion_id": "python_backend",
            "question_vi": "Bạn có thể mô tả một thay đổi backend do mình trực tiếp triển khai?",
            "purpose_vi": "Làm rõ phạm vi đóng góp kỹ thuật của ứng viên.",
            "source_span_ids": [ctx["span_id"]],
            "answer_indicators": ["Nêu quyết định kỹ thuật", "Giải thích cách xác minh kết quả"],
        }]
    }))
    async with test_session_factory() as session:
        await execute_interview_job(session, job_id=uuid.UUID(job_id), provider_override=mock_llm)
        await session.commit()

    async with AsyncClient(transport=transport, base_url="http://test", cookies=cookies, headers=headers) as client:
        detail = await client.get(f"/api/v1/interview-drafts/{draft_id}")
        assert detail.status_code == 200
        assert detail.json()["question_bank_id"] is None
        assert detail.json()["core_questions"] == []
        assert len(detail.json()["ai_followups"]) == 1
        assert detail.json()["is_stale"] is False


@pytest.mark.asyncio
async def test_interview_scorecard_saves_human_ratings_without_turning_missing_into_zero(test_session_factory, sample_docx_cv):
    async with test_session_factory() as session:
        ctx = await setup_interview_test_context(session, sample_docx_cv)

    transport = ASGITransport(app=app)
    owner_cookies = {SESSION_COOKIE_NAME: ctx["o_token"]}
    owner_headers = {"X-CSRF-Token": ctx["o_csrf"]}
    reviewer_cookies = {SESSION_COOKIE_NAME: ctx["r_token"]}
    reviewer_headers = {"X-CSRF-Token": ctx["r_csrf"]}
    criteria = [{
        "criterion_id": criterion.value,
        "outcome": "assessed" if criterion == CriterionId.PYTHON_BACKEND else "not_observed",
        "score": 0 if criterion == CriterionId.PYTHON_BACKEND else None,
        "answer_summary": "Không nêu được bước kiểm tra lỗi trong tình huống phỏng vấn." if criterion == CriterionId.PYTHON_BACKEND else "",
        "interviewer_note": "",
    } for criterion in CriterionId]

    async with AsyncClient(transport=transport, base_url="http://test", cookies=owner_cookies, headers=owner_headers) as client:
        invalid = [dict(row) for row in criteria]
        invalid[1]["score"] = 0
        invalid_res = await client.put(
            f"/api/v1/applications/{ctx['app_id']}/interview-scorecards",
            json={"round_no": 1, "expected_version": 0, "criteria": invalid},
        )
        assert invalid_res.status_code == 422

        saved = await client.put(
            f"/api/v1/applications/{ctx['app_id']}/interview-scorecards",
            json={"round_no": 1, "expected_version": 0, "criteria": criteria},
        )
        assert saved.status_code == 200
        card = saved.json()
        assert card["status"] == "draft"
        assert card["row_version"] == 1
        assert card["is_stale"] is False
        assert next(row for row in card["criteria"] if row["criterion_id"] == "python_backend")["score"] == 0
        assert next(row for row in card["criteria"] if row["criterion_id"] == "api_design")["score"] is None

        conflict = await client.put(
            f"/api/v1/applications/{ctx['app_id']}/interview-scorecards",
            json={"round_no": 1, "expected_version": 0, "criteria": criteria},
        )
        assert conflict.status_code == 409

    # A panel member can save their own draft without exposing it to the owner
    # before submission; the owner can still see their own private draft.
    async with test_session_factory() as session:
        session.add(InterviewScorecard(
            application_id=ctx["app_id"],
            interviewer_id=ctx["reviewer"].id,
            rubric_version_id=ctx["rubric_id"],
            round_no=1,
            status="draft",
            criteria_payload=criteria,
            source_snapshot={
                "application_generation": 1,
                "document_id": str(ctx["doc_id"]),
                "sanitized_version_id": str(ctx["sanitized_id"]),
                "rubric_version_id": str(ctx["rubric_id"]),
            },
            snapshot_hash="9" * 64,
            row_version=1,
        ))
        await session.commit()

    async with AsyncClient(transport=transport, base_url="http://test", cookies=owner_cookies, headers=owner_headers) as client:
        owner_list = await client.get(f"/api/v1/applications/{ctx['app_id']}/interview-scorecards")
        assert owner_list.status_code == 200
        assert [entry["id"] for entry in owner_list.json()] == [card["id"]]

    async with AsyncClient(transport=transport, base_url="http://test", cookies=reviewer_cookies, headers=reviewer_headers) as client:
        reviewer_list = await client.get(f"/api/v1/applications/{ctx['app_id']}/interview-scorecards")
        assert reviewer_list.status_code == 200
        assert len(reviewer_list.json()) == 1
        assert reviewer_list.json()[0]["interviewer_id"] == str(ctx["reviewer"].id)
        assert reviewer_list.json()[0]["id"] != card["id"]

    async with AsyncClient(transport=transport, base_url="http://test", cookies=owner_cookies, headers=owner_headers) as client:
        finalized = await client.post(
            f"/api/v1/interview-scorecards/{card['id']}/finalize",
            json={"expected_version": 1},
        )
        assert finalized.status_code == 200
        assert finalized.json()["status"] == "finalized"
        assert finalized.json()["row_version"] == 2

        locked = await client.put(
            f"/api/v1/applications/{ctx['app_id']}/interview-scorecards",
            json={"round_no": 1, "expected_version": 2, "criteria": criteria},
        )
        assert locked.status_code == 409


@pytest.mark.asyncio
async def test_interview_worker_blocks_revoked_cv_before_provider_call(test_session_factory, sample_docx_cv):
    async with test_session_factory() as session:
        ctx = await setup_interview_test_context(session, sample_docx_cv)

    async with AsyncClient(
        transport=ASGITransport(app=app),
        base_url="http://test",
        cookies={SESSION_COOKIE_NAME: ctx["o_token"]},
        headers={"X-CSRF-Token": ctx["o_csrf"]},
    ) as client:
        bank_res = await client.post(
            f"/api/v1/rubrics/{ctx['rubric_id']}/interview-question-banks",
            json={"source": "seed"},
        )
        bank_id = bank_res.json()["id"]
        approved = await client.post(
            f"/api/v1/interview-question-banks/{bank_id}/approve",
            json={"expected_rubric_version_id": str(ctx["rubric_id"]), "acknowledged": True},
        )
        assert approved.status_code == 200
        draft_res = await client.post(
            f"/api/v1/applications/{ctx['app_id']}/interview-drafts",
            json={
                "effective_result": {"kind": "assessment_run", "id": str(ctx["run_id"])},
                "expected_question_bank_id": bank_id,
            },
        )
        assert draft_res.status_code == 202

    async with test_session_factory() as session:
        version = await session.get(SanitizedVersion, ctx["sanitized_id"])
        version.status = SanitizedVersionStatus.REVOKED
        await session.commit()

    provider = MockLLMProvider()
    async with test_session_factory() as session:
        await execute_interview_job(session, uuid.UUID(draft_res.json()["job_id"]), provider_override=provider)
        await session.commit()
        draft = await session.get(InterviewDraft, uuid.UUID(draft_res.json()["id"]))
        assert draft.status == "failed"
    assert provider.invocation_count == 0

    async with AsyncClient(
        transport=ASGITransport(app=app),
        base_url="http://test",
        cookies={SESSION_COOKIE_NAME: ctx["o_token"]},
    ) as client:
        quarantined = await client.get(f"/api/v1/interview-drafts/{draft_res.json()['id']}")
        assert quarantined.status_code == 410
