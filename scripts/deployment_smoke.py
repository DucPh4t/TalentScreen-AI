"""Exercise the containerized HTTP workflow using synthetic data and mock LLM only.

Run against an isolated Compose project, never the real recruitment database.
Credentials file must contain SMOKE_LOGIN and SMOKE_PASSWORD; it is not logged.
This script creates a synthetic requisition/application and does not sign G1-G7.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timedelta, timezone
import io
import json
from pathlib import Path
import time
from urllib.parse import urlsplit

import docx
from dotenv import dotenv_values
import httpx


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-url", default="http://localhost:12004")
    parser.add_argument("--credentials-env", required=True)
    parser.add_argument("--output-metadata")
    args = parser.parse_args()
    if urlsplit(args.base_url).hostname not in {"localhost", "127.0.0.1"}:
        parser.error("Smoke checks are limited to an isolated loopback deployment.")
    credentials = dotenv_values(args.credentials_env)
    root = Path(__file__).resolve().parents[1]

    with httpx.Client(base_url=args.base_url, timeout=30) as client:
        def call(method: str, path: str, expected: int = 200, **kwargs):
            response = client.request(method, "/api/v1" + path, **kwargs)
            if response.status_code != expected:
                # Never echo response bodies, passwords, cookies, CVs or tokens.
                raise RuntimeError(f"{method} {path}: expected {expected}, got {response.status_code}")
            return response.json() if response.content else None

        health = call("GET", "/health")
        if health["app_env"] != "sandbox" or health["llm_provider"] != "mock":
            raise RuntimeError("Refusing to run outside sandbox + mock mode")
        call("GET", "/auth/me", 401)
        call("GET", "/config/diagnostic", 401)
        call("GET", "/admin/readiness")
        call("POST", "/auth/login", 401, json={"login_name": "not-a-user", "password": "incorrect"})
        login = call("POST", "/auth/login", json={
            "login_name": credentials["SMOKE_LOGIN"], "password": credentials["SMOKE_PASSWORD"],
        })
        user_id = login["user"]["id"]
        call("GET", "/auth/me")
        call("POST", "/requisitions", 403, json={"title": "Missing CSRF must fail"})
        client.headers["X-CSRF-Token"] = login["csrf_token"]
        print("PASS frontend proxy, readiness, login, anonymous access and CSRF")

        req = call("POST", "/requisitions", 201, json={"title": "SYNTHETIC Deployment Review — Backend Python"})
        req_id = req["id"]
        jd = call("POST", f"/requisitions/{req_id}/jd-versions", 201, json={
            "source_text": (root / "talentscreen-mvp-plan/examples/jd-backend-python.vi.md").read_text(),
            "change_reason": "Synthetic deployment smoke check",
            "expected_requisition_version": req["row_version"],
        })
        rubric = call("POST", f"/requisitions/{req_id}/rubrics", 201, json={"source": "seed"})
        req = call("GET", f"/requisitions/{req_id}")
        call("POST", f"/rubrics/{rubric['id']}/approve", json={
            "expected_requisition_version": req["row_version"],
            "expected_jd_version_id": jd["id"], "acknowledge_thresholds": True,
        })
        call("POST", f"/jd-versions/{jd['id']}/approve-egress", json={
            "expected_text_hash": jd["text_hash"], "acknowledged": True,
        })
        req = call("GET", f"/requisitions/{req_id}")
        call("PATCH", f"/requisitions/{req_id}", json={"status": "open", "expected_version": req["row_version"]})
        application = call("POST", f"/requisitions/{req_id}/applications", 201, json={})
        app_id = application["id"]
        document = docx.Document()
        document.add_heading("Synthetic technical experience", 0)
        document.add_paragraph("Developed Python backend services with FastAPI and PostgreSQL. Designed REST APIs with validation and pagination.")
        document.add_paragraph("Wrote pytest integration tests and debugged SQL query performance. Used Docker, CI/CD and structured logging.")
        document.add_paragraph("Implemented RBAC authorization, protected user data, and prepared deployment documentation.")
        content = io.BytesIO()
        document.save(content)
        upload = call("POST", f"/applications/{app_id}/documents", 202,
                      headers={"If-Match": str(application["row_version"])},
                      files={"file": ("synthetic-review.docx", content.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})

        def await_job(job_id: str) -> None:
            deadline = time.monotonic() + 120
            while time.monotonic() < deadline:
                job = call("GET", f"/jobs/{job_id}")
                if job["status"] == "succeeded":
                    return
                if job["status"] in {"failed", "cancelled", "superseded"}:
                    raise RuntimeError(f"Synthetic job ended: {job['status']}; code={job['last_error_code']}")
                time.sleep(1)
            raise RuntimeError("Synthetic job did not finish within 120 seconds")

        await_job(upload["job_id"])
        call("POST", f"/applications/{app_id}/raw-grants", 201, json={
            "grantee_user_id": user_id, "scopes": ["raw_cv"],
            "expires_at": (datetime.now(timezone.utc) + timedelta(minutes=10)).isoformat(),
            "reason": "Review synthetic sanitized text for deployment smoke check",
        })
        sanitized = call("GET", f"/documents/{upload['document']['id']}/sanitized-versions")
        source = call("GET", f"/sanitized-versions/{sanitized[-1]['id']}")
        application = call("GET", f"/applications/{app_id}")
        call("POST", f"/sanitized-versions/{source['id']}/approve", json={
            "expected_application_version": application["row_version"],
            "expected_sha256": source["sha256"], "acknowledged": True,
        })
        print("PASS JD, seed rubric, approval, intake, background worker and sanitization")

        assessment = call("POST", f"/applications/{app_id}/assessments", 202, json={
            "sanitized_version_id": source["id"], "rubric_version_id": rubric["id"],
        })
        deadline = time.monotonic() + 120
        while time.monotonic() < deadline:
            assessment = call("GET", f"/applications/{app_id}/assessments/{assessment['id']}")
            if assessment["status"] in {"succeeded", "completed"}:
                break
            if assessment["status"] == "failed":
                raise RuntimeError(f"Mock assessment failed: {assessment['failure_code']}")
            time.sleep(1)
        else:
            raise RuntimeError("Mock assessment did not finish within 120 seconds")
        assert len(assessment["criteria"]) == 6
        for criterion in assessment["criteria"]:
            if criterion["status"] == "assessed":
                assert criterion["evidence"], "Assessed score without evidence"
            else:
                assert criterion["score"] is None
            for evidence in criterion["evidence"]:
                span = call("GET", f"/source-spans/{evidence['span_id']}")
                assert span["text"] == evidence["quote"]
        progress = call("GET", f"/applications/{app_id}/progress")
        assert progress["stage"] == "awaiting_decision"
        copilot = call("POST", f"/applications/{app_id}/copilot", json={"message": "CV còn thiếu bằng chứng nào?"})
        assert copilot["provider"] == "mock" and copilot["facts"]
        assert all(fact["score"] is None for fact in copilot["facts"])
        if args.output_metadata:
            Path(args.output_metadata).write_text(json.dumps({"requisition_id": req_id, "application_id": app_id, "rubric_id": rubric["id"], "synthetic_only": True}, indent=2))
        call("GET", f"/requisitions/{req_id}/comparison")
        # Owner/admin workflow must not impersonate the independently assigned
        # reviewer. Reviewer context/submission is covered by isolated tests.
        call("GET", f"/applications/{app_id}/independent-review/context", 403)
        call("GET", f"/requisitions/{req_id}/independent-reviews")
        call("GET", "/admin/metrics")
        call("POST", "/auth/logout")
        call("GET", "/auth/me", 401)
        print("PASS mock assessment, grounded evidence, comparison, independent-review role isolation, metrics and logout")


if __name__ == "__main__":
    main()
